# Build flow-friendly Lesson 2 files (teacher v5 + student v2):
#  - NO page_break_before anywhere -> content flows continuously, user controls breaks
#  - "boxes" are shaded/bordered PARAGRAPHS (pBdr + w:shd), not 1x1 tables -> freely editable
#  - fixed column widths on 2-col tables; writing lines = bottom-border empty paragraphs
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = "/root/.hermes/cache/documents/doc_70fd9299bc97_Revised_Lesson_1_Relative_Clauses_Teacher's_CORRECTED.docx"
DIR = "/opt/f5-grammar-lesson2"
OUT_T = DIR + "/Grammar_Lesson_2_Inversions_Teacher_v5.docx"
OUT_S = DIR + "/Grammar_Lesson_2_Inversions_Student_v2.docx"

QUIZ_HDR, CORRECT, STRUCTURE, NOTEBOX, MODELBOX, INSTBOX = 'D9D9D9', 'E2EFDA', 'DEEBF7', 'F2F2F2', 'FFFFFF', 'E4DFEC'
GREY = RGBColor(0x59, 0x59, 0x59)

def new_doc():
    doc = docx.Document(SRC)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.15
    body = doc.element.body
    for child in list(body):
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)
    return doc

def style_run(r, font='Times New Roman', size=11, color=None):
    r.font.name = font; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color

def add_para(doc, text="", style=None, bold=False, italic=False, size=11, font='Times New Roman', color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text); r.bold, r.italic = bold, italic
        style_run(r, font, size, color)
    return p

def add_para_runs(doc, runs, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    for spec in runs:
        r = p.add_run(spec[0]); r.bold, r.italic = spec[1], spec[2]
        style_run(r)
    return p

def pbox(doc, runs, fill=None, color=None, borders=('top', 'left', 'bottom', 'right')):
    """Shaded/bordered PARAGRAPH box - normal text, freely editable/movable."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    pPr = p._p.get_or_add_pPr()
    if borders:
        pbdr = OxmlElement('w:pBdr')
        for edge in borders:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '4')
            pbdr.append(el)
        pPr.append(pbdr)
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
        pPr.append(shd)
    for spec in runs:
        r = p.add_run(spec[0]); r.bold, r.italic = spec[1], spec[2]
        style_run(r, color=color)
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_runs(cell, spec, color=None):
    cell.paragraphs[0].paragraph_format.line_spacing = 1.15
    for run in spec:
        r = cell.paragraphs[0].add_run(run[0])
        r.bold, r.italic = run[1], run[2]
        style_run(r, color=color)

def make_table(doc, rows, fills=None, widths_cm=None):
    fills = fills or {}
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles['Table Grid']
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    for i, row in enumerate(rows):
        for j, spec in enumerate(row):
            c = t.rows[i].cells[j]
            cell_runs(c, spec)
            if widths_cm: c.width = Cm(widths_cm[j])
            if (i, j) in fills: shade_cell(c, fills[(i, j)])
    return t

def write_line(doc):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '1')
    pbdr.append(el); pPr.append(pbdr)

def header(doc):
    for text, font, italic in [("St. Joseph’s Anglo-Chinese School", "Century Gothic", True),
                               ("F.5 English Language", "Trebuchet MS", True),
                               ("Grammar – Lesson 2: Inversions", "Trebuchet MS", True),
                               ("Inversions", "Times New Roman", False)]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text); r.bold = True; r.italic = italic; style_run(r, font=font)

R = lambda t: [(t, False, False)]
RB = lambda t: [(t, True, False)]

import lesson2_v4_content as C

W_STRUCT = [4.2, 12.4]
W_QUIZ = [7.0, 9.6]

def part1(doc, student=False):
    add_para(doc, "Part 1: Proofreading quiz", style="Heading 2")
    add_para(doc, "Each sentence contains ONE error. Underline the error and write the full corrected sentence. "
             "Items 1–5 review relative clauses from Lesson 1; Item 6 previews today’s topic — inversions.")
    rows = [[R("Sentence"), R("Correction" if student else "Correction (Teacher’s version)")]]
    def norm_cell(cell):
        if isinstance(cell, tuple):
            return [(s, False, False) for s in cell]
        return [tuple(run) for run in cell]
    for row in C.QUIZ_ROWS[1:]:
        sent = row[0][0] if isinstance(row[0], tuple) else row[0][0][0]
        answer = norm_cell(row[1]) if len(row) > 1 and not student else [("", False, False)]
        rows.append([R(sent), answer])
    fills = {(0, 0): QUIZ_HDR, (0, 1): QUIZ_HDR}
    if not student:
        for i in range(1, 7): fills[(i, 1)] = CORRECT
    make_table(doc, rows, fills=fills, widths_cm=W_QUIZ)

def part2(doc, student=False):
    add_para(doc, "Part 2: Three patterns for formal writing", style="Heading 2")
    add_para(doc, "The three patterns in this part all follow the same rule: the auxiliary verb (or the verb ‘be’) "
             "moves before the subject. Each pattern can lift a formal paragraph towards a higher band" +
             ("" if student else "; one more advanced pattern is kept in the teacher’s note at the end of this part") + ".")
    for title, structure, use, example, para_runs in C.PATTERNS:
        add_para(doc, title, style="Heading 3")
        make_table(doc, [
            [R("Structure:"), R(structure)],
            [R("Use in formal writing:"), R(use)],
            [R("Example:"), R(example)],
        ], fills={(i, j): STRUCTURE for i in range(3) for j in range(2)}, widths_cm=W_STRUCT)
        pbox(doc, para_runs)
    if not student:
        add_para(doc, C.NOTE_TITLE, style="Heading 3", color=GREY)
        add_para(doc, C.NOTE_INTRO, color=GREY)
        add_para(doc, C.NOTE_NAME, style="Heading 4", color=GREY)
        make_table(doc, [
            [R("Structure:"), R(C.NOTE_STRUCTURE)],
            [R("Example:"), R(C.NOTE_EXAMPLE)],
            [R("Use:"), R(C.NOTE_USE)],
        ], fills={(i, j): NOTEBOX for i in range(3) for j in range(2)}, widths_cm=W_STRUCT)
        add_para(doc, C.NOTE_EXTRA, color=GREY)

def part3(doc, student=False):
    add_para(doc, "Part 3: Practice sets", style="Heading 2")
    add_para(doc, "Rewrite each simple sentence using the pattern given in the prompt. The sentences are all about "
             "technology and communication, so keep the register formal.")
    for title, simple, prompt, answer in C.SETS:
        add_para(doc, title, style="Heading 3")
        fills = {(2, 1): CORRECT} if not student else {}
        rows = [
            [R("Simple sentence:"), R(simple)],
            [R("Prompt:"), [(prompt, False, True)]],
        ]
        if not student:
            rows.append([RB("Teacher’s Answer:"), R(answer)])
        make_table(doc, rows, fills=fills, widths_cm=W_STRUCT)

def part4(doc, student=False):
    add_para(doc, "Part 4: Writing Task", style="Heading 1")
    runs = [("Instructions:", True, False)]
    runs.append(("", False, False))  # placeholder replaced below with break
    p = pbox(doc, runs, fill=INSTBOX)
    # rebuild runs with a proper line break after label
    p._p.getparent().remove(p._p)
    runs2 = [("Instructions:", True, False)]
    runs2.append((C.INSTR_MAIN + ("" if student else "") + C.INSTR_BOLD, False, False))
    # simpler: plain shaded paragraph, break via add_break
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    pPr = p2._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '4')
        pbdr.append(el)
    pPr.append(pbdr)
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), INSTBOX)
    pPr.append(shd)
    r = p2.add_run("Instructions:"); r.bold = True; style_run(r)
    r2 = p2.add_run(); r2.add_break()
    r3 = p2.add_run(C.INSTR_MAIN + C.INSTR_BOLD); style_run(r3)
    if student:
        for _ in range(8):
            write_line(doc)
    else:
        pbox(doc, C.MODEL)
        pbox(doc, C.MODEL_NOTE, fill=NOTEBOX, color=GREY)

# ---------- build teacher v5 ----------
doc = new_doc(); header(doc)
part1(doc); part2(doc); part3(doc); part4(doc)
doc.save(OUT_T); print("SAVED:", OUT_T)

# ---------- build student v2 ----------
doc = new_doc(); header(doc)
part1(doc, student=True); part2(doc, student=True); part3(doc, student=True); part4(doc, student=True)
doc.save(OUT_S); print("SAVED:", OUT_S)