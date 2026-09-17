# Build Grammar Lesson 2 v2 (teacher's version): 2 core patterns + 2 optional in teacher's note.
# Colour-coded shaded boxes for readability. Reuses Lesson 1 docx as template for styles/page setup.
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = "/root/.hermes/cache/documents/doc_70fd9299bc97_Revised_Lesson_1_Relative_Clauses_Teacher's_CORRECTED.docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v2.docx"
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)

# palette
QUIZ_HDR  = 'D9D9D9'   # grey header
CORRECT   = 'E2EFDA'   # light green answers/corrections
STRUCTURE = 'DEEBF7'   # light blue structure boxes
MODELBOX  = 'FFF2CC'   # light yellow model paragraphs
NOTEBOX   = 'F2F2F2'   # grey teacher's notes / optional patterns
INSTBOX   = 'E4DFEC'   # light lavender instructions
GREY = RGBColor(0x59, 0x59, 0x59)

doc = docx.Document(SRC)
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

def add_para(text="", style=None, bold=False, italic=False, size=None, font=None,
             color=None, page_break_before=False):
    p = doc.add_paragraph(style=style)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
        if font: r.font.name = font
        if color: r.font.color.rgb = color
    return p

def add_para_runs(runs, style=None, page_break_before=False):
    """runs: list of (text, bold, italic[, size, font, color])"""
    p = doc.add_paragraph(style=style)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    for spec in runs:
        r = p.add_run(spec[0])
        r.bold, r.italic = spec[1], spec[2]
        if len(spec) > 3 and spec[3]: r.font.size = Pt(spec[3])
        if len(spec) > 4 and spec[4]: r.font.name = spec[4]
        if len(spec) > 5 and spec[5]: r.font.color.rgb = spec[5]
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_runs(cell, spec, color=None):
    p = cell.paragraphs[0]
    for run in spec:
        r = p.add_run(run[0])
        r.bold, r.italic = run[1], run[2]
        if len(run) > 3 and run[3]: r.font.size = Pt(run[3])
        if color: r.font.color.rgb = color

def make_table(rows, fills=None):
    """rows: list of rows; each row = list of cell specs (list of run tuples).
    fills: dict {(i,j): hex} -> shade that cell."""
    fills = fills or {}
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles['Table Grid']
    for i, row in enumerate(rows):
        for j, spec in enumerate(row):
            c = t.rows[i].cells[j]
            cell_runs(c, spec)
            if (i, j) in fills:
                shade_cell(c, fills[(i, j)])
    return t

def box(shading, runs, color=None):
    """Single-cell shaded box (paragraph-like shape with border)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles['Table Grid']
    c = t.rows[0].cells[0]
    shade_cell(c, shading)
    cell_runs(c, runs, color=color)
    return t

R  = lambda t: [(t, False, False)]
RB = lambda t: [(t, True, False)]

# ============ HEADER (same as Lesson 1) ============
add_para_runs([("St. Joseph’s Anglo-Chinese School", True, True, None, "Century Gothic")])
add_para_runs([("F.5 English Language", True, True, None, "Trebuchet MS")])
add_para_runs([("Grammar – Lesson 2: Inversions", True, True, None, "Trebuchet MS")])
add_para_runs([("Inversions", True, False, None, "Times New Roman")])

# ============ PART 1: PROOFREADING QUIZ ============
add_para("Part 1: Proofreading quiz", style="Heading 2")
add_para("Each sentence contains ONE error. Underline the error and write the full corrected sentence. "
         "Items 1–5 review relative clauses from Lesson 1; Item 6 previews today’s topic — inversions.")
quiz_rows = [
    [R("Sentence"), R("Correction (Teacher’s version)")],
    [R("Social media connects people instantly which makes the world feel smaller."),
     R("Social media connects people instantly, which makes the world feel smaller. "
       "(Add a comma before the connective clause.)")],
    [R("Online courses are ideal for busy adults, especially those which have little time to attend classes."),
     R("Online courses are ideal for busy adults, especially those who have little time to attend classes. "
       "(People take \u2018who\u2019, not \u2018which\u2019.)")],
    [R("There are dozens of messaging apps, all of them are free to use."),
     R("There are dozens of messaging apps, all of which are free to use. "
       "(A pronoun cannot join two clauses; use \u2018which\u2019.)")],
    [R("The forum which students share their homework is very popular."),
     [("Teacher’s Answer A: The forum on which students share their homework is very popular. ", True, False),
      ("Teacher’s Answer B: The forum where students share their homework is very popular. "
       "(Add the missing preposition \u2018on\u2019.)", False, False)]],
    [R("Many teenagers spend hours gaming, that worries their parents."),
     R("Many teenagers spend hours gaming, which worries their parents. "
       "(\u2018That\u2019 cannot begin a connective clause.)")],
    [R("Never the school has seen such strong demand for online lessons."),
     R("Never has the school seen such strong demand for online lessons. "
       "(Move the auxiliary before the subject \u2014 more on this today.)")],
]
fills_quiz = {(0, 0): QUIZ_HDR, (0, 1): QUIZ_HDR}
for i in range(1, 7):
    fills_quiz[(i, 1)] = CORRECT
make_table(quiz_rows, fills=fills_quiz)

# ============ PART 2: TWO CORE PATTERNS ============
add_para("Part 2: Two patterns for formal writing", style="Heading 2", page_break_before=True)
add_para("Both patterns follow the same rule: the auxiliary verb (or the verb \u2018be\u2019) moves before the "
         "subject. We focus on two patterns that are high-value and safe for formal writing; two more advanced "
         "patterns are kept in the teacher’s note at the end of this part.")

# Pattern 1
add_para("Pattern 1: Not only ... but also ...", style="Heading 3")
make_table([
    [R("Structure:"), R("Not only + auxiliary + subject + verb, but (subject) also + verb")],
    [R("Use in formal writing:"), R("Adds a second advantage in proposals, reports and articles.")],
    [R("Example:"), R("Not only does the scheme save money, but it also builds confidence.")],
], fills={(0, 0): STRUCTURE, (0, 1): STRUCTURE, (1, 0): STRUCTURE, (1, 1): STRUCTURE,
          (2, 0): STRUCTURE, (2, 1): STRUCTURE})
box(MODELBOX, [
    ("Introducing a school e-learning platform would bring clear benefits. ", False, False),
    ("Not only does the platform give students round-the-clock access to learning materials, but it also allows teachers to track progress more closely", False, True),
    (". With such a system, revision becomes a habit rather than a scramble before exams.", False, False)])

# Pattern 2
add_para("Pattern 2: Never / Rarely / Seldom", style="Heading 3")
make_table([
    [R("Structure:"), R("Never / Rarely / Seldom + auxiliary + subject + verb")],
    [R("Use in formal writing:"), R("Turns a criticism or a surprising claim into a formal, striking opening.")],
    [R("Example:"), R("Rarely do we question our daily screen habits.")],
], fills={(0, 0): STRUCTURE, (0, 1): STRUCTURE, (1, 0): STRUCTURE, (1, 1): STRUCTURE,
          (2, 0): STRUCTURE, (2, 1): STRUCTURE})
box(MODELBOX, [
    ("Many parents blame video games for every poor result. ", False, False),
    ("Rarely do they stop to consider how limited their children’s other entertainment options are", False, True),
    (". The real problem may be the lack of balance, not the games themselves.", False, False)])

# ---- Additional patterns: teacher's note (grey) ----
add_para("Additional patterns (teacher’s note — optional)", style="Heading 3", color=GREY)
add_para("Introduce these two only after the class is secure with Patterns 1 and 2. The mechanism is the same: "
         "front the expression, then place the auxiliary before the subject.", color=GREY)
add_para("Only + adverbial (only when / only after / only by)", style="Heading 4", color=GREY)
make_table([
    [R("Structure:"), R("Only + when / after / by + clause or phrase + auxiliary + subject + verb")],
    [R("Example:"), R("Only when students manage their screen time can they learn effectively.")],
], fills={(0, 0): NOTEBOX, (0, 1): NOTEBOX, (1, 0): NOTEBOX, (1, 1): NOTEBOX})
add_para("Trap: \u2018Only + noun\u2019 does NOT invert. \u2018Only students can join\u2019 is already correct; "
         "inversion happens only when \u2018only\u2019 is followed by an adverbial.", color=GREY)
add_para("So ... that / Such ... that", style="Heading 4", color=GREY)
make_table([
    [R("Structure:"), R("So + adjective + auxiliary + subject + verb + that ...  OR  Such + be + noun phrase + that ...")],
    [R("Example:"), R("Such is the influence of social media that it shapes public opinion overnight.")],
], fills={(0, 0): NOTEBOX, (0, 1): NOTEBOX, (1, 0): NOTEBOX, (1, 1): NOTEBOX})
add_para("Use it to open with a strong statement when emphasising the scale or importance of an issue.", color=GREY)

# ============ PART 3: PRACTICE SETS ============
add_para("Part 3: Practice sets", style="Heading 2", page_break_before=True)
add_para("Rewrite each simple sentence using the pattern given in the prompt. The sentences are all about "
         "technology and communication, so keep the register formal.")

sets = [
    ("Example Set 1 (Pattern 1 – Not only ... but also ...)",
     "The online platform saves paper, and it gives students instant feedback.",
     "Begin with \u2018Not only\u2019 to add a second advantage.",
     "Not only does the online platform save paper, but it also gives students instant feedback."),
    ("Example Set 2 (Pattern 1 – Not only ... but also ...)",
     "E-books are cheaper than printed textbooks, and they are far lighter to carry.",
     "Use \u2018Not only ... but also ...\u2019 to combine the two advantages of e-books.",
     "Not only are e-books cheaper than printed textbooks, but they are also far lighter to carry."),
    ("Example Set 3 (Pattern 1 – Not only ... but also ...)",
     "The new platform provides video lessons, and it lets students take quizzes at home.",
     "Begin with \u2018Not only\u2019 to add the second feature.",
     "Not only does the new platform provide video lessons, but it also lets students take quizzes at home."),
    ("Example Set 4 (Pattern 2 – Never / Rarely / Seldom)",
     "People rarely question the amount of time they spend scrolling on their phones.",
     "Begin with \u2018Rarely\u2019.",
     "Rarely do people question the amount of time they spend scrolling on their phones."),
    ("Example Set 5 (Pattern 2 – Never / Rarely / Seldom)",
     "People rarely consider how much personal data they share online.",
     "Begin with \u2018Rarely\u2019.",
     "Rarely do people consider how much personal data they share online."),
    ("Example Set 6 (Pattern 2 – Never / Rarely / Seldom)",
     "The workplace has never relied so heavily on technology.",
     "Begin with \u2018Never\u2019.",
     "Never has the workplace relied so heavily on technology."),
]
for title, simple, prompt, answer in sets:
    add_para(title, style="Heading 3")
    make_table([
        [R("Simple sentence:"), R(simple)],
        [R("Prompt:"), [(prompt, False, True)]],
        [RB("Teacher’s Answer:"), R(answer)],
    ], fills={(2, 1): CORRECT})

# ============ PART 4: WRITING TASK ============
add_para("Part 4: Writing Task", style="Heading 1", page_break_before=True)
ti = doc.add_table(rows=1, cols=1); ti.style = doc.styles['Table Grid']
ci = ti.rows[0].cells[0]
shade_cell(ci, INSTBOX)
r = ci.paragraphs[0].add_run("Instructions:"); r.bold = True
r2 = ci.paragraphs[0].add_run(); r2.add_break()
ci.paragraphs[0].add_run("Social media plays a big part in students’ lives. Write ONE paragraph (about 80–100 words) "
                         "giving your view on whether students should limit their use of social media. "
                         "Use at least two inversions and one relative clause. ")
r3 = ci.paragraphs[0].add_run("Remember: the flow of ideas comes first — never force in an inversion just for the sake of using one.")
r3.bold = True

model_runs = [
    ("Many people accuse social media of destroying students’ concentration and turning them into passive consumers of short videos. ", False, False, 10.5),
    ("Rarely do they mention the benefits it brings to learning, such as instant access to discussion groups and subject resources,", True, True, 10.5),
    (" (inversion) ", False, False, 10.5),
    ("which older generations could never enjoy", True, True, 10.5),
    (" (relative clause). ", False, False, 10.5),
    ("Not only does social media connect students with their classmates, but it also links them to news and ideas from around the world.", True, True, 10.5),
    (" (inversion). ", False, False, 10.5),
    ("The wiser response, then, is not to ban it but to teach young people how to use it well.", False, False, 10.5),
]
box(MODELBOX, model_runs)
box(NOTEBOX, [
    ("Teacher’s note — flow before form: ", True, False),
    ("In the model, every inversion earns its place. \u2018Rarely\u2019 contrasts directly with the accusation in "
     "sentence 1, and \u2018Not only ... but also ...\u2019 adds the second advantage with a natural rhythm. "
     "No inversion is added for decoration. Ask students to read their paragraphs aloud: if an inversion makes "
     "the flow awkward, they should re-order the ideas instead of forcing the structure.", False, False)])

doc.save(OUT)
print("SAVED:", OUT)