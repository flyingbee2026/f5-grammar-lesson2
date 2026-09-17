# Build Grammar Lesson 2 v4 (teacher's version): 3 core patterns, 6 practice sets,
# writing task = ONE inversion + ONE relative clause. TNR 11, spacing 1.15.
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = "/root/.hermes/cache/documents/doc_70fd9299bc97_Revised_Lesson_1_Relative_Clauses_Teacher's_CORRECTED.docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v4.docx"
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)

QUIZ_HDR, CORRECT, STRUCTURE, NOTEBOX, INSTBOX = 'D9D9D9', 'E2EFDA', 'DEEBF7', 'F2F2F2', 'E4DFEC'
GREY = RGBColor(0x59, 0x59, 0x59)

doc = docx.Document(SRC)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.paragraph_format.line_spacing = 1.15
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

def style_run(r, font='Times New Roman', size=11, color=None):
    r.font.name = font; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color

def add_para(text="", style=None, bold=False, italic=False, size=11, font='Times New Roman',
             color=None, page_break_before=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    if page_break_before: p.paragraph_format.page_break_before = True
    if text:
        r = p.add_run(text); r.bold, r.italic = bold, italic
        style_run(r, font, size, color)
    return p

def add_para_runs(runs, style=None, page_break_before=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    if page_break_before: p.paragraph_format.page_break_before = True
    for spec in runs:
        r = p.add_run(spec[0]); r.bold, r.italic = spec[1], spec[2]
        style_run(r)
    return p

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_runs(cell, spec, color=None):
    cell.paragraphs[0].paragraph_format.line_spacing = 1.15
    for run in spec:
        r = cell.paragraphs[0].add_run(run[0])
        r.bold, r.italic = run[1], run[2]
        style_run(r, color=color)

def make_table(rows, fills=None):
    fills = fills or {}
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles['Table Grid']
    for i, row in enumerate(rows):
        for j, spec in enumerate(row):
            c = t.rows[i].cells[j]
            cell_runs(c, spec)
            if (i, j) in fills: shade_cell(c, fills[(i, j)])
    return t

def box(runs, color=None):
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles['Table Grid']
    cell_runs(t.rows[0].cells[0], runs, color=color)
    return t

R = lambda t: [(t, False, False)]
RB = lambda t: [(t, True, False)]

import lesson2_v4_content as C

# ============ HEADER (same as Lesson 1) ============
for text, font, italic in [("St. Joseph’s Anglo-Chinese School", "Century Gothic", True),
                           ("F.5 English Language", "Trebuchet MS", True),
                           ("Grammar – Lesson 2: Inversions", "Trebuchet MS", True),
                           ("Inversions", "Times New Roman", False)]:
    p = doc.add_paragraph(style="Normal"); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); r.bold = True; r.italic = italic; style_run(r, font=font)

# ============ PART 1: PROOFREADING QUIZ ============
add_para("Part 1: Proofreading quiz", style="Heading 2")
add_para("Each sentence contains ONE error. Underline the error and write the full corrected sentence. "
         "Items 1–5 review relative clauses from Lesson 1; Item 6 previews today’s topic — inversions.")
quiz_rows = [[R(c[0]) if len(c) == 1 else c for c in row] for row in C.QUIZ_ROWS]
fills_quiz = {(0, 0): QUIZ_HDR, (0, 1): QUIZ_HDR}
for i in range(1, 7): fills_quiz[(i, 1)] = CORRECT
make_table(quiz_rows, fills=fills_quiz)

# ============ PART 2: THREE CORE PATTERNS ============
add_para("Part 2: Three patterns for formal writing", style="Heading 2", page_break_before=True)
add_para("The three patterns in this part all follow the same rule: the auxiliary verb (or the verb ‘be’) moves "
         "before the subject. Each pattern can lift a formal paragraph towards a higher band; one more advanced "
         "pattern is kept in the teacher’s note at the end of this part.")
for title, structure, use, example, para_runs in C.PATTERNS:
    add_para(title, style="Heading 3")
    make_table([
        [R("Structure:"), R(structure)],
        [R("Use in formal writing:"), R(use)],
        [R("Example:"), R(example)],
    ], fills={(i, j): STRUCTURE for i in range(3) for j in range(2)})
    box(para_runs)

# ---- Additional pattern: teacher's note (grey) ----
add_para(C.NOTE_TITLE, style="Heading 3", color=GREY)
add_para(C.NOTE_INTRO, color=GREY)
add_para(C.NOTE_NAME, style="Heading 4", color=GREY)
make_table([
    [R("Structure:"), R(C.NOTE_STRUCTURE)],
    [R("Example:"), R(C.NOTE_EXAMPLE)],
    [R("Use:"), R(C.NOTE_USE)],
], fills={(i, j): NOTEBOX for i in range(3) for j in range(2)})
add_para(C.NOTE_EXTRA, color=GREY)

# ============ PART 3: PRACTICE SETS ============
add_para("Part 3: Practice sets", style="Heading 2", page_break_before=True)
add_para("Rewrite each simple sentence using the pattern given in the prompt. The sentences are all about "
         "technology and communication, so keep the register formal.")
for title, simple, prompt, answer in C.SETS:
    add_para(title, style="Heading 3")
    make_table([
        [R("Simple sentence:"), R(simple)],
        [R("Prompt:"), [(prompt, False, True)]],
        [RB("Teacher’s Answer:"), R(answer)],
    ], fills={(2, 1): CORRECT})

# ============ PART 4: WRITING TASK ============
add_para("Part 4: Writing Task", style="Heading 1", page_break_before=True)
ti = doc.add_table(rows=1, cols=1); ti.style = doc.styles['Table Grid']
ci = ti.rows[0].cells[0]; shade_cell(ci, INSTBOX)
ci.paragraphs[0].paragraph_format.line_spacing = 1.15
r = ci.paragraphs[0].add_run("Instructions:"); r.bold = True; style_run(r)
r2 = ci.paragraphs[0].add_run(); r2.add_break()
r3 = ci.paragraphs[0].add_run(C.INSTR_MAIN); style_run(r3)
r4 = ci.paragraphs[0].add_run(C.INSTR_BOLD); r4.bold = True; style_run(r4)

box(C.MODEL)
box(C.MODEL_NOTE, color=GREY)

doc.save(OUT)
print("SAVED:", OUT)