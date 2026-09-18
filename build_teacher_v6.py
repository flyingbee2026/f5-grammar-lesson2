# Build Grammar Lesson 2 TEACHER v6 mirroring Helen's revised student version:
#  - header 'Grammar – Lesson 1: Inversion' (her title), borders under header lines 1-3
#  - quiz: 6x1 table, TWO errors per sentence + grey teacher's answers table
#  - part 2 intro without the teacher-note mention; Pattern 2 use = her wording
#  - sets: 3-row tables (Simple/Prompt/Teacher's Answer), no fills
#  - part 4: lavender box with numbered instructions; model in grey box (no writing lines)
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = "/root/.hermes/cache/documents/doc_70fd9299bc97_Revised_Lesson_1_Relative_Clauses_Teacher's_CORRECTED.docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v6.docx"

STRUCTURE, NOTEBOX, INSTBOX, ANSBOX = 'DEEBF7', 'F2F2F2', 'E4DFEC', 'F2F2F2'
GREY = RGBColor(0x59, 0x59, 0x59)

doc = docx.Document(SRC)
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.paragraph_format.line_spacing = 1.15
body = doc.element.body
for child in list(body):
    if child.tag in (qn('w:p'), qn('w:tbl')):
        body.remove(child)

def style_run(r, font='Times New Roman', size=11, color=None):
    r.font.name = font; r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color

def add_para(text="", style=None, bold=False, italic=False, size=11, font='Times New Roman', color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text); r.bold, r.italic = bold, italic
        style_run(r, font, size, color)
    return p

def pbox(runs, fill=None, color=None, borders=('top', 'left', 'bottom', 'right')):
    """Bordered/shaded paragraph box (normal, editable text)."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    pPr = p._p.get_or_add_pPr()
    if borders:
        pbdr = OxmlElement('w:pBdr')
        for edge in borders:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '4')
            pbdr.append(el)
        pPr.append(pbdr)
    if fill:
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), fill)
        pPr.append(shd)
    for spec in runs:
        r = p.add_run(spec[0]); r.bold, r.italic = spec[1], spec[2]
        style_run(r, color=color)
    return p

def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:space'), '4')
    pbdr.append(el); pPr.append(pbdr)

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def cell_runs(cell, spec, color=None):
    cell.paragraphs[0].paragraph_format.line_spacing = 1.15
    for run in spec:
        r = cell.paragraphs[0].add_run(run[0])
        r.bold, r.italic = run[1], run[2]
        style_run(r, color=color)

def make_table(rows, fills=None):
    fills = fills or {}
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles['Table Grid']
    for i, row in enumerate(rows):
        for j, spec in enumerate(row):
            c = t.rows[i].cells[j]
            cell_runs(c, spec)
            if (i, j) in fills: shade_cell(c, fills[(i, j)])
    return t

R = lambda t: [(t, False, False)]
RB = lambda t: [(t, True, False)]

# ============ HEADER (mirrors her student file) ============
for text, font, italic, border in [("St. Joseph’s Anglo-Chinese School", "Century Gothic", True, True),
                                   ("F.5 English Language", "Trebuchet MS", True, True),
                                   ("Grammar – Lesson 1: Inversion", "Trebuchet MS", True, True),
                                   ("Inversion", "Times New Roman", False, False)]:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); r.bold = True; r.italic = italic; style_run(r, font=font)
    if border: bottom_border(p)

# ============ PART 1: PROOFREADING QUIZ (TWO errors each) ============
add_para("Part 1: Proofreading quiz", style="Heading 2")
add_para("Each sentence contains TWO errors. Proofread the sentences.")
quiz = [
    "Social media connects people instantly which make the world feel smaller.",
    "Online courses are ideal for busy adults, especially those which has little time to attend classes.",
    "There are dozens of messaging apps, all of them is free to use.",
    "The forums which students share their homework is very popular.",
    "Many teenagers spend hours gaming, that worry their parents.",
    "Not only the internet connects people around the world, but it also spreads information quickly.",
]
make_table([[R(s)] for s in quiz])

# teacher's answers table
answers = [
    ("1", "Social media connects people instantly, which makes the world feel smaller. "
          "(Add the comma; make \u2192 makes \u2014 the clause refers to the whole idea.)"),
    ("2", "Online courses are ideal for busy adults, especially those who have little time to attend classes. "
          "(which \u2192 who; has \u2192 have.)"),
    ("3", "There are dozens of messaging apps, all of which are free to use. "
          "(them \u2192 which; is \u2192 are.)"),
    ("4", "Teacher’s Answer A: The forums on which students share their homework are very popular. "
          "Teacher’s Answer B: The forums where students share their homework are very popular. "
          "(Add on; is \u2192 are.)"),
    ("5", "Many teenagers spend hours gaming, which worries their parents. "
          "(that \u2192 which \u2014 connective clause; worry \u2192 worries.)"),
    ("6", "Not only does the internet connect people around the world, but it also spreads information quickly. "
          "(Add does before the subject; connects \u2192 connect.)"),
]
add_para("")
add_para("Teacher’s answers", bold=True)
rows = [[R("Item"), RB("Correction (fix both errors)")]]
for num, ans in answers:
    rows.append([R(num), R(ans)])
fills = {(i, j): ANSBOX for i in range(len(rows)) for j in range(2)}
make_table(rows, fills=fills)

# ============ PART 2: THREE CORE PATTERNS ============
add_para("Part 2: Three patterns for formal writing", style="Heading 2")
add_para("The three patterns in this part all follow the same rule: the auxiliary verb (or the verb ‘be’) moves "
         "before the subject. Each pattern can lift a formal paragraph towards a higher band.")

patterns = [
    ("Pattern 1: Not only ... but also ...",
     "Not only + auxiliary + subject + verb, but (subject) also + verb",
     "Adds a second advantage in proposals, reports and articles.",
     "Not only does the scheme save money, but it also builds confidence.",
     [("Introducing a school e-learning platform would bring clear benefits. ", False, False),
      ("Not only does the platform give students round-the-clock access to learning materials, but it also allows teachers to track progress more closely", False, False),
      (". With such a system, revision becomes a habit rather than a scramble before exams.", False, False)]),
    ("Pattern 2: Never / Rarely / Seldom",
     "Never / Rarely / Seldom + auxiliary + subject + verb",
     "Makes a criticism or a claim sound stronger and more formal.",
     "Rarely do we question our daily screen habits.",
     [("Many parents blame video games for every poor result. ", False, False),
      ("Rarely do they stop to consider how limited their children’s other entertainment options are", False, False),
      (". The real problem may be the lack of balance, not the games themselves.", False, False)]),
    ("Pattern 3: Only by ...",
     "Only by + -ing / noun phrase + auxiliary + subject + verb",
     "Sets a condition that leads to a result — ideal for conclusions and recommendations.",
     "Only by upgrading the campus network can the school offer reliable online lessons.",
     [("Working from home has become common in many industries. ", False, False),
      ("Only by investing in secure networks and clear communication tools can companies keep their teams productive", False, False),
      (". Technology can help, but good management matters more.", False, False)]),
]
for title, structure, use, example, para_runs in patterns:
    add_para(title, style="Heading 3")
    make_table([
        [R("Structure:"), R(structure)],
        [R("Use in formal writing:"), R(use)],
        [R("Example:"), R(example)],
    ], fills={(i, j): STRUCTURE for i in range(3) for j in range(2)})
    pbox(para_runs)

# optional pattern note (teacher only)
add_para("Additional pattern (teacher’s note — optional)", style="Heading 3", color=GREY)
add_para("Introduce this only after the class is secure with the three core patterns. The mechanism is the same: "
         "front the expression, then place the auxiliary before the subject.", color=GREY)
add_para("So ... that / Such ... that", style="Heading 4", color=GREY)
make_table([
    [R("Structure:"), R("So + adjective + auxiliary + subject + verb + that ...  OR  Such + be + noun phrase + that ...")],
    [R("Example:"), R("Such is the influence of social media that it shapes public opinion overnight.")],
    [R("Use:"), R("Use it to open with a strong statement when emphasising the scale or importance of an issue.")],
], fills={(i, j): NOTEBOX for i in range(3) for j in range(2)})
add_para("You can also extend Pattern 3 with ‘Only when ...’ or ‘Only after ...’ for stronger classes.", color=GREY)

# ============ PART 3: PRACTICE SETS ============
add_para("Part 3: Practice sets", style="Heading 2")
add_para("Rewrite each simple sentence using the pattern given in the prompt. The sentences are all about "
         "technology and communication, so keep the register formal.")
sets = [
    ("Example Set 1 (Not only ... but also ...)",
     "The online platform saves paper, and it gives students instant feedback.",
     "Begin with ‘Not only’ to add a second advantage.",
     "Not only does the online platform save paper, but it also gives students instant feedback."),
    ("Example Set 2 (Not only ... but also ...)",
     "E-books are cheaper than printed textbooks, and they are far lighter to carry.",
     "Use ‘Not only ... but also ...’ to combine the two advantages of e-books.",
     "Not only are e-books cheaper than printed textbooks, but they are also far lighter to carry."),
    ("Example Set 3 (Never / Rarely / Seldom)",
     "Parents rarely check what their children are doing online.",
     "Begin with ‘Rarely’.",
     "Rarely do parents check what their children are doing online."),
    ("Example Set 4 (Never / Rarely / Seldom)",
     "The workplace has never relied so heavily on technology.",
     "Begin with ‘Never’.",
     "Never has the workplace relied so heavily on technology."),
    ("Example Set 5 (Only by ...)",
     "Schools can protect students’ privacy by issuing clear guidelines on social media use.",
     "Begin with ‘Only by’.",
     "Only by issuing clear guidelines on social media use can schools protect students’ privacy."),
    ("Example Set 6 (Only by ...)",
     "Bookshops can survive in the digital age by offering experiences that apps cannot match.",
     "Begin with ‘Only by’.",
     "Only by offering experiences that apps cannot match can bookshops survive in the digital age."),
]
for title, simple, prompt, answer in sets:
    add_para(title, style="Heading 3")
    make_table([
        [R("Simple sentence:"), R(simple)],
        [R("Prompt:"), [(prompt, False, True)]],
        [RB("Teacher’s Answer:"), R(answer)],
    ])

# ============ PART 4: WRITING TASK ============
add_para("Part 4: Writing Task", style="Heading 1")
for runs in [([("Instructions:", True, False)]),
             ([("1. Social media plays a big part in students’ lives. Write ONE paragraph (about 80–100 "
                "words) giving your view on whether students should limit their use of social media. ", False, False)]),
             ([("2. Use ONE inversion and ONE relative clause. ", False, False)])]:
    pbox(runs, fill=INSTBOX)
pbox([("Teacher’s model answer: ", True, False),
      ("Social media is often blamed for distracting students, and every few months someone suggests banning it "
       "from school. Banning, however, treats the symptom, not the cause. Rarely do students have a more "
       "convenient way to exchange ideas and share learning resources, which textbooks alone can never provide. "
       "Through discussion groups, they can prepare for exams together and correct each other’s mistakes in real "
       "time. The real task, therefore, is to teach young people to use social media responsibly, not to lock it "
       "away.", False, False)], fill=NOTEBOX)

doc.save(OUT)
print("SAVED:", OUT)