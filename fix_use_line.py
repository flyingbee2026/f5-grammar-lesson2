import docx

NEW = "Highlights how rare or unusual something is — makes a criticism or a claim sound stronger and more formal."
OLD = "Turns a criticism or a surprising claim into a formal, striking opening."

# patch content module
path = "/opt/f5-grammar-lesson2/lesson2_v4_content.py"
src = open(path, encoding="utf-8").read()
assert OLD in src, "OLD not found in content module"
open(path, "w", encoding="utf-8").write(src.replace(OLD, NEW))
print("content module patched")

def patch_table(docpath, table_idx, row_idx, col_idx):
    d = docx.Document(docpath)
    cell = d.tables[table_idx].rows[row_idx].cells[col_idx]
    p = cell.paragraphs[0]
    # check it's the target cell
    found = False
    for r in p.runs:
        if "Turns a criticism" in r.text:
            found = True
    if not found:
        # maybe text split across runs; concat check
        if "Turns a criticism" in p.text:
            found = True
    if not found:
        print("WARN: target not in", docpath, "table", table_idx, "->", p.text[:60])
        return False
    # replace: clear runs, add single run with same base formatting
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(NEW)
    from docx.shared import Pt
    r.font.name = "Times New Roman"; r.font.size = Pt(11)
    d.save(docpath)
    print("patched:", docpath)
    return True

patch_table("/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v5.docx", 2, 1, 1)
patch_table("/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Student_v2.docx", 2, 1, 1)