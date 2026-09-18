# Patch Helen's latest STUDENT docx -> Student v4: Set 4 prompt -> 'Begin with Seldom'
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = "/root/.hermes/cache/documents/doc_4bc471b74964_Lesson 2 - Inversion (Student's).docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Student_v4.docx"

d = docx.Document(SRC)

def set_cell_text(cell, text):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)

done = False
for t in d.tables:
    if len(t.rows) == 3 and len(t.columns) == 2 \
            and t.rows[0].cells[1].text.strip().startswith('Teenagers often go to sleep'):
        set_cell_text(t.rows[1].cells[1], "Begin with ‘Seldom’.")
        done = True
        break
print('set 4 prompt updated:', done)
d.save(OUT)
print('SAVED:', OUT)