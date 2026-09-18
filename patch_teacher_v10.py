# Patch Helen's attached TEACHER docx -> Teacher v10:
#  - Set 4: 'Begin with Seldom' + Seldom answer
#  - Part 4: replace single model with THREE sample answers (one per pattern, each + relative clause)
import copy
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = "/root/.hermes/cache/documents/doc_91f688413127_Lesson 2 - Inversion (Teacher's).docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v10.docx"

d = docx.Document(SRC)

def set_text(p, text, bold=False, size=11):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    return r

def set_cell(cell, runs):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, bold in runs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = 'Times New Roman'; r.font.size = Pt(11)

NEW_SETS = {
    'Example Set 2 (Not only ... but also ...)': {
        'simple': "The workshop improved students’ digital skills, and it boosted their confidence in online presentations.",
        'prompt': "Use ‘Not only ... but also ...’ to combine the workshop’s two benefits.",
        'answer': "Not only did the workshop improve students’ digital skills, but it also boosted their confidence in online presentations."},
    'Example Set 3 (Never / Rarely / Seldom)': {
        'simple': "Readers do not often verify news sources before sharing them.",
        'prompt': "Begin with ‘Rarely’.",
        'answer': "Rarely do readers verify news sources before sharing them."},
    'Example Set 4 (Never / Rarely / Seldom)': {
        'simple': "Teenagers often go to sleep without turning off their notifications, which leads to hours of late-night scrolling.",
        'prompt': "Begin with ‘Seldom’.",
        'answer': "Seldom do teenagers turn off their notifications before going to sleep, which leads to hours of late-night scrolling."},
    'Example Set 5 (Only by ...)': {
        'simple': "Schools can protect students’ privacy if they issue clear guidelines on social media use.",
        'prompt': "Begin with ‘Only by’.",
        'answer': "Only by issuing clear guidelines on social media use can schools protect students’ privacy."},
    'Example Set 6 (Only by ...)': {
        'simple': "If we want to fight online scams, governments should work closely with technology companies.",
        'prompt': "Begin with ‘Only by’.",
        'answer': "Only by working closely with technology companies can governments fight online scams."},
}

SAMPLES = [
    ("Sample answer 1 (Not only ... but also ...)",
     "People often blame social media for distracting students, yet the same platforms bring real learning "
     "benefits. Not only does social media give students instant access to discussion groups and subject "
     "resources, but it also keeps them connected with classmates, which makes group revision much easier. "
     "Limiting use, rather than banning it, seems the wiser policy."),
    ("Sample answer 2 (Never / Rarely / Seldom)",
     "Many adults simply assume that social media wastes students’ time. Rarely do they consider how it allows "
     "young people to follow the news, learn new skills and build online portfolios, which are all valued in "
     "further studies and future work. The sensible answer, therefore, is to guide students’ use rather than "
     "to condemn it."),
    ("Sample answer 3 (Only by ...)",
     "Social media is undoubtedly attractive, and students will not give it up on their own. Only by setting "
     "sensible limits and talking openly about online habits can parents help their children, who often struggle "
     "to control screen time themselves. When parents understand the pressures behind late-night scrolling, a "
     "supportive conversation beats a sudden ban."),
]

def fill_sample_el(el, label, text):
    """el: CT_P element; clear runs, add bold label run + text run (TNR 11)."""
    for r in el.findall(qn('w:r')):
        el.remove(r)
    p = Paragraph(el, d)
    r1 = p.add_run(label + ' ')
    r1.bold = True; r1.font.name = 'Times New Roman'; r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.name = 'Times New Roman'; r2.font.size = Pt(11)

body = d.element.body
last_heading = None
patched = 0
for child in list(body):
    if child.tag == qn('w:p'):
        p = Paragraph(child, d)
        if p.style.name == 'Heading 3' and p.text.strip():
            last_heading = p.text.strip()
    elif child.tag == qn('w:tbl'):
        t = Table(child, d)
        rows, cols = len(t.rows), len(t.columns)
        if rows == 3 and cols == 2 and t.rows[0].cells[0].text.strip() == 'Structure:' \
                and 'builds confidence' in t.rows[2].cells[1].text:
            set_cell(t.rows[2].cells[1], [("Not only does the scheme save money, but it also builds students’ confidence.", False)])
        if rows == 3 and cols == 2 and t.rows[0].cells[0].text.strip().startswith('Simple sentence') \
                and last_heading in NEW_SETS:
            info = NEW_SETS[last_heading]
            set_cell(t.rows[0].cells[1], [(info['simple'], False)])
            set_cell(t.rows[1].cells[1], [(info['prompt'], False)])
            set_cell(t.rows[2].cells[1], [(info['answer'], False)])
            patched += 1
            print('SET patched:', last_heading)

for p in list(d.paragraphs):
    if p.style.name == 'Normal' and p.text.strip().startswith('Rewrite each simple sentence'):
        set_text(p, p.text.rstrip() + ' The even-numbered sets (2, 4 and 6) are more challenging.')
    if p.style.name == 'Normal' and 'about 80–100 words' in p.text:
        set_text(p, p.text.replace('about 80–100 words', 'about 60 words'))
    if p.style.name == 'Normal' and p.text.strip().startswith('Teacher’s model answer'):
        # transform this paragraph into sample 1, then clone twice for samples 2 and 3
        fill_sample_el(p._p, SAMPLES[0][0], SAMPLES[0][1])
        cur = p._p
        for label, text in SAMPLES[1:]:
            # deepcopy the paragraph element (keeps pBdr + shading), then clear its runs
            new_el = copy.deepcopy(cur)
            for r in new_el.findall(qn('w:r')):
                new_el.remove(r)
            cur.addnext(new_el)
            fill_sample_el(new_el, label, text)
            cur = new_el
        print('3 sample answers written')

print('sets patched:', patched)
d.save(OUT)
print('SAVED:', OUT)