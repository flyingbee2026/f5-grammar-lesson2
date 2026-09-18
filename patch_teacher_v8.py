# Patch Helen's attached TEACHER docx -> Teacher v8, aligning it with her revised STUDENT version.
import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = "/root/.hermes/cache/documents/doc_91f688413127_Lesson 2 - Inversion (Teacher's).docx"
OUT = "/opt/f5-grammar-lesson2/Grammar_Lesson_2_Inversions_Teacher_v8.docx"

d = docx.Document(SRC)

def set_text(p, text, bold=False, size=11):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    return r

def set_cell(cell, runs):
    """runs: list of (text, bold). Rebuild paragraph 0 of the cell."""
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, bold in runs:
        r = p.add_run(text)
        r.bold = bold
        r.font.name = 'Times New Roman'; r.font.size = Pt(11)

NEW_SETS = {
    'Example Set 2 (Not only ... but also ...)': {
        'simple': "The workshop improved students’ digital skills, and it boosted their confidence in online presentations.",
        'prompt': "Use ‘Not only ... but also ...’ to combine the workshop’s two benefits.",
        'answer': "Not only did the workshop improve students’ digital skills, but it also boosted their confidence in online presentations."},
    'Example Set 3 (Never / Rarely / Seldom)': {
        'simple': "Parents don’t really check what their children are doing online.",
        'prompt': "Begin with ‘Rarely’.",
        'answer': "Rarely do parents really check what their children are doing online."},
    'Example Set 4 (Never / Rarely / Seldom)': {
        'simple': "People rarely question headline claims, and they should not trust unverified rumors.",
        'prompt': "Begin with ‘Never’.",
        'answer': "Never should people trust unverified rumors, and rarely do they question headline claims. "
                  "(Alternative: People rarely question headline claims, and never should they trust unverified rumors.)"},
    'Example Set 5 (Only by ...)': {
        'simple': "Schools can protect students’ privacy if they issue clear guidelines on social media use.",
        'prompt': "Begin with ‘Only by’.",
        'answer': "Only by issuing clear guidelines on social media use can schools protect students’ privacy."},
    'Example Set 6 (Only by ...)': {
        'simple': "If we want to fight online scams, governments should work closely with technology companies.",
        'prompt': "Begin with ‘Only by’.",
        'answer': "Only by working closely with technology companies can governments fight online scams."},
}

body = d.element.body
last_heading = None
patched = 0
for child in list(body):
    if child.tag == qn('w:p'):
        p = Paragraph(child, d)
        if p.style.name == 'Heading 3' and p.text.strip():
            last_heading = p.text.strip()
    elif child.tag == qn('w:tbl'):
        t = Table(child, d)
        rows, cols = len(t.rows), len(t.columns)
        # pattern-1 example clarity
        if rows == 3 and cols == 2 and t.rows[0].cells[0].text.strip() == 'Structure:' \
                and 'builds confidence' in t.rows[2].cells[1].text:
            set_cell(t.rows[2].cells[1], [("Not only does the scheme save money, but it also builds students’ confidence.", False)])
            print('P1 example ->', t.rows[2].cells[1].text[:60])
        # sets
        if rows == 3 and cols == 2 and t.rows[0].cells[0].text.strip().startswith('Simple sentence') \
                and last_heading in NEW_SETS:
            info = NEW_SETS[last_heading]
            set_cell(t.rows[0].cells[1], [(info['simple'], False)])
            set_cell(t.rows[1].cells[1], [(info['prompt'], False)])
            set_cell(t.rows[2].cells[1], [(info['answer'], False)])
            patched += 1
            print('SET patched:', last_heading)

# part 3 intro
for p in list(d.paragraphs):
    if p.style.name == 'Normal' and p.text.strip().startswith('Rewrite each simple sentence'):
        set_text(p, p.text.rstrip() + ' The even-numbered sets (2, 4 and 6) are more challenging.')
        print('part3 intro updated')
    # writing task instruction: 80-100 -> 60 words
    if p.style.name == 'Normal' and 'about 80–100 words' in p.text:
        set_text(p, p.text.replace('about 80–100 words', 'about 60 words'))
        print('instruction updated:', p.text[:80])
    # model answer -> trimmed to ~60 words
    if p.style.name == 'Normal' and p.text.strip().startswith('Teacher’s model answer'):
        model = ("Social media is often blamed for distracting students, and every few months someone suggests "
                 "banning it. A ban, however, treats the symptom, not the cause. Rarely do students have a more "
                 "convenient way to exchange ideas and share learning resources, which textbooks alone can never "
                 "provide. The real task is to teach responsible use, not to lock the apps away.")
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        r1 = p.add_run('Teacher’s model answer: '); r1.bold = True
        r1.font.name = 'Times New Roman'; r1.font.size = Pt(11)
        r2 = p.add_run(model); r2.font.name = 'Times New Roman'; r2.font.size = Pt(11)
        print('model updated (%d words)' % len(model.split()))

print('sets patched:', patched)
d.save(OUT)
print('SAVED:', OUT)